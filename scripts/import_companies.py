#!/usr/bin/env python3
"""Turns newly-added rows in onboarding/companies.xlsx into a real,
working company — the same files a developer would otherwise hand-write
(see docs/client-onboarding.md's manual path, and every src/config/
companies/*.ts file, all of which were originally written this exact way).

What it does, per pending company row (status blank/"new"):
  1. Validates the row + its child rows (GMC Plans, Super Top-Up, Term
     Life, Addons) + that the logo file exists in
     onboarding/incoming/<id>/<logo_filename>. On any problem, writes a
     specific "error: ..." message into that row's status cell and moves
     on to the next company — one bad row never blocks the others.
  2. Copies the logo + the shared loophealth.png into
     public/branding/<id>/, and generates a QR code for the portal URL.
  3. Writes src/config/companies/<id>.ts (spreads ...tarentoProfile, same
     structure/comments as every hand-authored profile).
  4. Appends the new profile to src/config/companies/index.ts and to
     supabase/functions/export-plays/index.ts's companyTable map.
  5. Writes a new Supabase migration from supabase/MIGRATION_TEMPLATE.sql.txt.
  6. Marks the row "imported" and saves the workbook.

This produces a plain git diff for review — nothing here touches a live
database or deploys anything. Apply the generated migration and redeploy
export-plays yourself (see the printed summary) once the PR lands.

Usage:
    python3 scripts/import_companies.py
"""
from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import qrcode

REPO_ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = REPO_ROOT / "onboarding" / "companies.xlsx"
INCOMING_DIR = REPO_ROOT / "onboarding" / "incoming"
COMPANIES_DIR = REPO_ROOT / "src" / "config" / "companies"
BRANDING_DIR = REPO_ROOT / "public" / "branding"
MIGRATIONS_DIR = REPO_ROOT / "supabase" / "migrations"
MIGRATION_TEMPLATE_PATH = REPO_ROOT / "supabase" / "MIGRATION_TEMPLATE.sql.txt"
EXPORT_PLAYS_PATH = REPO_ROOT / "supabase" / "functions" / "export-plays" / "index.ts"
INDEX_TS_PATH = COMPANIES_DIR / "index.ts"
SKIP_STATUSES = {"imported", "example"}  # rows the importer never touches

ELIGIBILITY_MAP = {
    "everyone": ("m => true", "Any covered member"),
    "female adults 18+": ("m => m.gender==='female' && m.age>=18", "Female members, 18+"),
    "parents only": ("m => m.relation==='parent'", "Parents only"),
}


def slugify(raw: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", raw.strip().lower()).strip("_")
    return s[:32]


def js_str(value) -> str:
    """Renders a TS string literal, preferring single quotes to match the
    rest of the codebase's style, falling back to a JSON (double-quoted)
    literal — still valid TS, and correctly escaped — whenever the value
    itself contains a quote/backslash/newline that would break that."""
    s = "" if value is None else str(value)
    if "'" not in s and "\\" not in s and "\n" not in s:
        return f"'{s}'"
    return json.dumps(s, ensure_ascii=False)


def js_value(value) -> str:
    """Like js_str, but passes through None as the TS literal `null`
    instead of an empty string — for fields that are meaningfully
    optional (e.g. a plan's `note`), not just "a string that happens to
    be empty"."""
    return "null" if value is None else js_str(value)


def to_int(value, default=0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def sheet_rows(ws) -> list[dict]:
    headers = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        rows.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
    return rows


@dataclass
class Issue:
    row_idx: int
    message: str


@dataclass
class CompanyInput:
    row_idx: int
    id: str
    company_name: str
    tagline: str
    portal_url: str
    max_spouse: int
    max_child: int
    max_parent: int
    logo_filename: str
    gmc_plans: list = field(default_factory=list)
    super_top_ups: list = field(default_factory=list)
    gtl_plans: list = field(default_factory=list)
    addons: list = field(default_factory=list)


def company_profile_files():
    for p in COMPANIES_DIR.glob("*.ts"):
        if p.stem in ("index", "types", "_template"):
            continue
        yield p


def load_company_inputs(wb) -> tuple[list[CompanyInput], list[Issue], object]:
    company_ws = wb["Company"]
    headers = [c.value for c in company_ws[1]]
    status_col = headers.index("status") + 1

    pending: list[CompanyInput] = []
    issues: list[Issue] = []

    plans_by_id = {}
    for r in sheet_rows(wb["GMC Plans"]):
        plans_by_id.setdefault(r.get("company_id"), []).append(r)
    topups_by_id = {}
    for r in sheet_rows(wb["Super Top-Up"]):
        if r.get("company_id"):
            topups_by_id.setdefault(r["company_id"], []).append(r)
    gtl_by_id = {}
    for r in sheet_rows(wb["Term Life"]):
        if r.get("company_id"):
            gtl_by_id.setdefault(r["company_id"], []).append(r)
    addons_by_id = {}
    for r in sheet_rows(wb["Addons"]):
        if r.get("company_id"):
            addons_by_id.setdefault(r["company_id"], []).append(r)

    for row_idx in range(2, company_ws.max_row + 1):
        row = {headers[i]: company_ws.cell(row=row_idx, column=i + 1).value for i in range(len(headers))}
        if all(v is None or str(v).strip() == "" for v in row.values()):
            continue
        status = str(row.get("status") or "").strip().lower()
        if status in SKIP_STATUSES or status.startswith("error"):
            continue

        raw_id = str(row.get("id") or "").strip()
        company_name = str(row.get("company_name") or "").strip()
        cid = slugify(raw_id) if raw_id else slugify(company_name)

        if not company_name:
            issues.append(Issue(row_idx, "error: company_name is required"))
            continue
        if not re.match(r"^[a-z][a-z0-9_]{1,31}$", cid):
            issues.append(Issue(row_idx, f"error: couldn't derive a valid id from '{raw_id or company_name}'"))
            continue

        pending.append(CompanyInput(
            row_idx=row_idx,
            id=cid,
            company_name=company_name,
            tagline=str(row.get("tagline") or "Benefits Adventure").strip(),
            portal_url=str(row.get("portal_url") or "https://benefits.loophealth.com/").strip(),
            max_spouse=to_int(row.get("max_spouse"), 1),
            max_child=to_int(row.get("max_child"), 2),
            max_parent=to_int(row.get("max_parent"), 2),
            logo_filename=str(row.get("logo_filename") or "").strip(),
            gmc_plans=plans_by_id.get(raw_id, []) or plans_by_id.get(cid, []),
            super_top_ups=topups_by_id.get(raw_id, []) or topups_by_id.get(cid, []),
            gtl_plans=gtl_by_id.get(raw_id, []) or gtl_by_id.get(cid, []),
            addons=addons_by_id.get(raw_id, []) or addons_by_id.get(cid, []),
        ))

    return pending, issues, status_col


def validate(ci: CompanyInput, seen_ids: set[str], existing_ids: set[str]) -> str | None:
    if ci.id in existing_ids:
        return f"error: id '{ci.id}' is already used by an existing company"
    if ci.id in seen_ids:
        return f"error: id '{ci.id}' is used by more than one row in this sheet"
    if not ci.gmc_plans or not any(str(p.get("plan_name") or "").strip() for p in ci.gmc_plans):
        return f"error: no GMC Plans rows found with company_id='{ci.id}' (check it matches the Company sheet's id exactly)"
    if not ci.logo_filename:
        return "error: logo_filename is required"
    logo_path = INCOMING_DIR / ci.id / ci.logo_filename
    if not logo_path.exists():
        return f"error: logo not found at onboarding/incoming/{ci.id}/{ci.logo_filename}"
    if logo_path.suffix.lower() not in (".png", ".jpg", ".jpeg", ".webp"):
        return f"error: logo must be .png, .jpg or .webp (got {logo_path.suffix})"
    return None


def build_gmc_plan_ts(idx: int, row: dict, dependant_note: str) -> str:
    name = str(row.get("plan_name") or f"Plan {idx+1}").strip()
    desc = str(row.get("description") or "").strip()
    si = to_int(row.get("sum_insured"))
    prem = to_int(row.get("premium"))
    features = []
    for i in range(1, 5):
        label = row.get(f"feature_{i}_label")
        value = row.get(f"feature_{i}_value")
        if label or value:
            features.append(f"{{ label: {js_str(label)}, value: {js_str(value)} }}")
    note = "Company Sponsored" if prem == 0 else None
    if features:
        features_lines = ",\n      ".join(features)
        features_block = f"[\n      {features_lines}\n    ]"
    else:
        features_block = "[]"
    return f"""  {{
    id: 'gmc_{idx+1}',
    name: {js_str(name)},
    desc: {js_str(desc)},
    coverNote: {js_str(dependant_note)},
    sums: [{{ si: {si}, prem: {prem}, note: {js_value(note)}, emi: null }}],
    features: {features_block},
  }},"""


def build_topup_or_gtl_ts(idx: int, row: dict, name_key: str, for_gmc_ids: list[str] | None) -> str:
    name = str(row.get(name_key) or f"Tier {idx+1}").strip()
    desc = str(row.get("description") or "").strip()
    si = to_int(row.get("sum_insured"))
    prem = to_int(row.get("premium"))
    if for_gmc_ids is None:
        return f"  {{ id: 'tier_{idx+1}', name: {js_str(name)}, desc: {js_str(desc)}, si: {si}, prem: {prem} }},"
    entries = []
    for gid in for_gmc_ids:
        entries.append(
            f"  {{ id: 'stu_{idx+1}_{gid}', forGmcPlanId: '{gid}', name: {js_str(name)}, "
            f"desc: {js_str(desc)}, si: {si}, prem: {prem}, emi: null }},"
        )
    return "\n".join(entries)


def build_addon_ts(idx: int, row: dict, company_id: str) -> str:
    name = str(row.get("addon_name") or f"Addon {idx+1}").strip()
    addon_id = f"{company_id}_addon_{idx+1}"
    price = to_int(row.get("price"))
    per = str(row.get("per") or "member").strip().lower()
    per = per if per in ("member", "family") else "member"
    elig_key = str(row.get("eligibility") or "Everyone").strip().lower()
    elig_fn, elig_note = ELIGIBILITY_MAP.get(elig_key, ELIGIBILITY_MAP["everyone"])
    desc = str(row.get("description") or "").strip()
    image_filename = str(row.get("image_filename") or "").strip()
    image_line = ""
    if image_filename:
        image_line = f"\n    imageUrl: assetPath('branding/addons/{addon_id}{Path(image_filename).suffix}'),"
    return f"""  {{
    id: {js_str(addon_id)}, name: {js_str(name)}, per: {js_str(per)}, price: {price},
    desc: {js_str(desc)},
    elig: {elig_fn}, eligNote: {js_str(elig_note)},
    brand: {js_str(name)},{image_line}
  }},"""


def generate_profile_ts(ci: CompanyInput, logo_ext: str) -> str:
    dependant_note = (
        f"Covers Self, up to {ci.max_spouse} Spouse, {ci.max_child} Children "
        f"and {ci.max_parent} Parent(s)/Parent(s)-in-law, per this company's policy."
    )
    gmc_rows = [r for r in ci.gmc_plans if str(r.get("plan_name") or "").strip()]
    gmc_plans_ts = "\n".join(build_gmc_plan_ts(i, r, dependant_note) for i, r in enumerate(gmc_rows))
    gmc_ids = [f"gmc_{i+1}" for i in range(len(gmc_rows))]

    lines = [
        "import { assetPath } from '../assetPath';",
        "import { tarentoProfile } from './tarento';",
        "import type { CompanyProfile } from './types';",
        "",
        "// Generated by scripts/import_companies.py from onboarding/companies.xlsx —",
        "// safe to hand-edit afterwards, this file won't be overwritten automatically.",
        "",
        f"const {ci.id}GmcPlans: CompanyProfile['gmcPlans'] = [",
        gmc_plans_ts,
        "];",
        "",
    ]

    has_topups = bool(ci.super_top_ups and any(str(r.get("name") or "").strip() for r in ci.super_top_ups))
    if has_topups:
        topup_rows = [r for r in ci.super_top_ups if str(r.get("name") or "").strip()]
        topup_ts = "\n".join(build_topup_or_gtl_ts(i, r, "name", gmc_ids) for i, r in enumerate(topup_rows))
        lines += [
            f"const {ci.id}SuperTopUpPlans: CompanyProfile['superTopUpPlans'] = [",
            topup_ts,
            "];",
            "",
        ]

    has_gtl = bool(ci.gtl_plans and any(str(r.get("tier_name") or "").strip() for r in ci.gtl_plans))
    if has_gtl:
        gtl_rows = [r for r in ci.gtl_plans if str(r.get("tier_name") or "").strip()]
        gtl_ts = "\n".join(build_topup_or_gtl_ts(i, r, "tier_name", None) for i, r in enumerate(gtl_rows))
        lines += [
            f"const {ci.id}GtlPlans: CompanyProfile['gtlPlans'] = [",
            gtl_ts,
            "];",
            "",
        ]

    has_addons = bool(ci.addons and any(str(r.get("addon_name") or "").strip() for r in ci.addons))
    if has_addons:
        addon_rows = [r for r in ci.addons if str(r.get("addon_name") or "").strip()]
        addons_ts = "\n".join(build_addon_ts(i, r, ci.id) for i, r in enumerate(addon_rows))
        lines += [
            f"const {ci.id}Addons: CompanyProfile['addons'] = [",
            addons_ts,
            "];",
            "",
        ]

    try:
        host = ci.portal_url.split("//", 1)[-1].split("/", 1)[0]
    except Exception:
        host = "benefits.loophealth.com"

    lines += [
        f"export const {ci.id}Profile: CompanyProfile = {{",
        "  ...tarentoProfile,",
        f"  id: {js_str(ci.id)},",
        f"  titleSuffix: {js_str(ci.company_name)},",
        f"  COMPANY: {js_str(ci.company_name)},",
        f"  TAGLINE: {js_str(ci.tagline)},",
        f"  LOGO_URL: assetPath('branding/{ci.id}/logo{logo_ext}'),",
        f"  LOGO_FALLBACK_URL: assetPath('branding/{ci.id}/logo{logo_ext}'),",
        "  partners: [",
        f"    {{ name: {js_str(ci.company_name)}, logoUrl: assetPath('branding/{ci.id}/logo{logo_ext}') }},",
        f"    {{ name: 'Loophealth', logoUrl: assetPath('branding/{ci.id}/loophealth.png') }},",
        "  ],",
        f"  PORTAL_URL: {js_str(ci.portal_url)},",
        f"  PORTAL_HOST_LABEL: {js_str(host)},",
        f"  PORTAL_QR_URL: assetPath('branding/{ci.id}/qr_enrol_portal.png'),",
        f"  PORTAL_SUPPORT_COPY: {js_str(ci.company_name + ' profile — complete enrolment on your company portal after selecting your coverage.')},",
        f"  dependantStructure: {{ spouse: {ci.max_spouse}, child: {ci.max_child}, parent: {ci.max_parent} }},",
        f"  gmcPlans: {ci.id}GmcPlans,",
        (f"  superTopUpPlans: {ci.id}SuperTopUpPlans," if has_topups
         else "  // No Super Top-Up rows in the sheet for this company — inherited from tarentoProfile (empty array)."),
        (f"  gtlPlans: {ci.id}GtlPlans," if has_gtl else "  gtlPlans: [],"),
        (f"  addons: {ci.id}Addons," if has_addons
         else "  // No custom Addons rows in the sheet — inherited from tarentoProfile (the standard catalogue)."),
        f"  enableLevel4: {'true' if has_gtl else 'false'},",
        f"  PLAYS_TABLE: {js_str('plays_' + ci.id)},",
        f"  HOF_RPC: {js_str('hall_of_fame_' + ci.id)},",
        "};",
        "",
    ]
    return "\n".join(lines)


def next_migration_timestamp(offset_seconds: int) -> str:
    existing = sorted(p.name[:14] for p in MIGRATIONS_DIR.glob("*.sql") if p.name[:14].isdigit())
    base = datetime.utcnow()
    if existing:
        latest = datetime.strptime(existing[-1], "%Y%m%d%H%M%S")
        base = max(base, latest)
    return (base + timedelta(seconds=offset_seconds + 1)).strftime("%Y%m%d%H%M%S")


def write_migration(ci: CompanyInput, offset_seconds: int) -> Path:
    template = MIGRATION_TEMPLATE_PATH.read_text()
    sql = template.replace("__SLUG__", ci.id)
    ts = next_migration_timestamp(offset_seconds)
    out_path = MIGRATIONS_DIR / f"{ts}_add_company_{ci.id}.sql"
    out_path.write_text(sql)
    return out_path


def register_in_index_ts(ci: CompanyInput):
    text = INDEX_TS_PATH.read_text()
    import_line = f"import {{ {ci.id}Profile }} from './{ci.id}';\n"
    last_import = list(re.finditer(r"^import .+\n", text, re.MULTILINE))[-1]
    text = text[:last_import.end()] + import_line + text[last_import.end():]
    # Append as the new last entry, right before the array's closing `];` —
    # matches how every company has been registered by hand so far, and
    # keeps the diff additive/easy to review instead of reordering the
    # existing list.
    text = re.sub(
        r"(export const companyProfiles: CompanyProfile\[\] = \[\n(?:.*\n)*?)(\];)",
        r"\1  " + f"{ci.id}Profile,\n" + r"\2",
        text,
        count=1,
    )
    INDEX_TS_PATH.write_text(text)


def register_in_export_plays(ci: CompanyInput):
    text = EXPORT_PLAYS_PATH.read_text()
    new_entry = f'  {ci.id}: "plays_{ci.id}",\n'
    text = re.sub(
        r"(const companyTable: Record<string, string> = \{\n(?:.*\n)*?)(\};)",
        r"\1" + new_entry + r"\2",
        text,
        count=1,
    )
    EXPORT_PLAYS_PATH.write_text(text)


def place_branding(ci: CompanyInput) -> str:
    out_dir = BRANDING_DIR / ci.id
    out_dir.mkdir(parents=True, exist_ok=True)
    src_logo = INCOMING_DIR / ci.id / ci.logo_filename
    ext = src_logo.suffix.lower()
    shutil.copyfile(src_logo, out_dir / f"logo{ext}")

    shared_loophealth = BRANDING_DIR / "loophealth.png"
    if shared_loophealth.exists():
        shutil.copyfile(shared_loophealth, out_dir / "loophealth.png")

    qr_img = qrcode.make(ci.portal_url)
    qr_img.save(out_dir / "qr_enrol_portal.png")

    custom_addon_images_dir = INCOMING_DIR / ci.id / "addons"
    if custom_addon_images_dir.is_dir():
        for i, row in enumerate(a for a in ci.addons if str(a.get("addon_name") or "").strip()):
            fname = str(row.get("image_filename") or "").strip()
            if not fname:
                continue
            src = custom_addon_images_dir / fname
            if src.exists():
                addon_id = f"{ci.id}_addon_{i+1}"
                shutil.copyfile(src, BRANDING_DIR / "addons" / f"{addon_id}{src.suffix.lower()}")
    return ext


def main():
    if not WORKBOOK_PATH.exists():
        print(f"No workbook found at {WORKBOOK_PATH}. Run scripts/build_onboarding_template.py first.")
        return

    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    pending, upfront_issues, status_col = load_company_inputs(wb)
    company_ws = wb["Company"]

    for issue in upfront_issues:
        company_ws.cell(row=issue.row_idx, column=status_col, value=issue.message)

    existing_ids = {p.stem for p in company_profile_files()}
    seen_ids: set[str] = set()
    imported, failed = [], []

    for offset, ci in enumerate(pending):
        err = validate(ci, seen_ids, existing_ids)
        if err:
            company_ws.cell(row=ci.row_idx, column=status_col, value=err)
            failed.append((ci.id or f"row {ci.row_idx}", err))
            continue
        seen_ids.add(ci.id)

        try:
            logo_ext = place_branding(ci)
            profile_ts = generate_profile_ts(ci, logo_ext)
            (COMPANIES_DIR / f"{ci.id}.ts").write_text(profile_ts)
            register_in_index_ts(ci)
            register_in_export_plays(ci)
            migration_path = write_migration(ci, offset)
        except Exception as e:  # noqa: BLE001 - surface any failure into the sheet, never crash the whole run
            company_ws.cell(row=ci.row_idx, column=status_col, value=f"error: {e}")
            failed.append((ci.id, str(e)))
            continue

        company_ws.cell(row=ci.row_idx, column=status_col, value="imported")
        existing_ids.add(ci.id)
        imported.append((ci, migration_path))

    wb.save(WORKBOOK_PATH)

    print(f"\n{len(imported)} company(ies) imported, {len(failed)} failed, "
          f"{len(upfront_issues)} row(s) had upfront errors.\n")
    for ci, migration_path in imported:
        print(f"✔ {ci.company_name} ({ci.id})")
        print(f"   src/config/companies/{ci.id}.ts")
        print(f"   public/branding/{ci.id}/")
        print(f"   {migration_path.relative_to(REPO_ROOT)}")
        print(f"   → still needed: apply that migration, then redeploy export-plays\n")
    for cid, msg in failed:
        print(f"✘ {cid}: {msg}")
    for issue in upfront_issues:
        print(f"✘ row {issue.row_idx}: {issue.message}")


if __name__ == "__main__":
    main()
